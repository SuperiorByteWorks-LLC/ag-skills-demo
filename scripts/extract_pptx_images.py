#!/usr/bin/env python3
"""
Extract all images from PowerPoint files and organize them for course materials.

Usage:
    python extract_pptx_images.py <pptx_file> [output_dir]
    python extract_pptx_images.py "3. Class Title.pptx" docs/classes/images

Output:
    Saves images with naming pattern: class##_slide##_img##.ext
    Example: class03_slide05_img01.png
"""

import os
import sys
from pathlib import Path
from pptx import Presentation
from PIL import Image
from io import BytesIO


def extract_class_number(pptx_path):
    """
    Extract class number from PowerPoint filename.
    
    Examples:
        "0. Welcome Class.pptx" -> "00"
        "3. Navigating the US Agricultural Data Landscape + Workshop.pptx" -> "03"
        "14. Future Farm.pptx" -> "14"
    """
    filename = Path(pptx_path).stem
    
    # Get first part before the dot
    first_part = filename.split('.')[0].strip()
    
    try:
        class_num = int(first_part)
        return f"{class_num:02d}"
    except ValueError:
        print(f"Warning: Could not extract class number from '{filename}'")
        return "XX"


def extract_images(pptx_path, output_dir="docs/classes/images"):
    """
    Extract all images from PowerPoint presentation.
    
    Args:
        pptx_path (str): Path to .pptx file
        output_dir (str): Directory to save images (default: docs/classes/images)
    
    Returns:
        tuple: (total_images_extracted, class_number, slides_processed)
    """
    # Verify input file exists
    pptx_file = Path(pptx_path)
    if not pptx_file.exists():
        print(f"Error: File not found: {pptx_path}")
        return 0, None, 0
    
    if not pptx_file.suffix.lower() == '.pptx':
        print(f"Error: File must be .pptx format, got {pptx_file.suffix}")
        return 0, None, 0
    
    # Create output directory
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    
    # Extract class number from filename
    class_num = extract_class_number(pptx_path)
    
    print(f"\nExtracting images from: {pptx_file.name}")
    print(f"Class number: {class_num}")
    print(f"Output directory: {output_path}")
    print("-" * 60)
    
    # Open presentation
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        print(f"Error opening PowerPoint file: {e}")
        return 0, class_num, 0
    
    total_images = 0
    slides_with_images = 0
    
    # Process each slide
    for slide_idx, slide in enumerate(prs.slides, 1):
        images_in_slide = 0
        
        # Process each shape in slide
        for shape_idx, shape in enumerate(slide.shapes):
            # Check if shape contains an image
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    # Extract image
                    image = shape.image
                    image_bytes = image.blob
                    extension = image.ext
                    
                    # Create filename
                    filename = f"class{class_num}_slide{slide_idx:02d}_img{images_in_slide:02d}.{extension}"
                    filepath = output_path / filename
                    
                    # Save image
                    with open(filepath, 'wb') as f:
                        f.write(image_bytes)
                    
                    # Verify image is valid by trying to open it
                    try:
                        with Image.open(filepath) as img:
                            img_size = img.size
                            print(f"  ✓ {filename:<45} ({img_size[0]}x{img_size[1]} px)")
                    except Exception as e:
                        print(f"  ! {filename:<45} (extracted but may be corrupted)")
                    
                    total_images += 1
                    images_in_slide += 1
                    
                except Exception as e:
                    print(f"  ✗ Error extracting image from slide {slide_idx}: {e}")
        
        # Track slides with images
        if images_in_slide > 0:
            slides_with_images += 1
    
    # Summary
    print("-" * 60)
    print(f"\nExtraction Summary:")
    print(f"  Total images extracted: {total_images}")
    print(f"  Slides with images: {slides_with_images} / {len(prs.slides)}")
    print(f"  Output location: {output_path}")
    print()
    
    return total_images, class_num, len(prs.slides)


def batch_extract(directory=".", output_dir="docs/classes/images"):
    """
    Extract images from all PowerPoint files in a directory.
    
    Args:
        directory (str): Directory containing .pptx files
        output_dir (str): Directory to save all images
    """
    directory = Path(directory)
    pptx_files = sorted(directory.glob("*.pptx"))
    
    if not pptx_files:
        print(f"No .pptx files found in {directory}")
        return
    
    print(f"\nBatch extracting from {len(pptx_files)} PowerPoint files...")
    print(f"Output directory: {output_dir}")
    print("=" * 60)
    
    total_all_images = 0
    total_all_slides = 0
    
    for pptx_file in pptx_files:
        images, class_num, slides = extract_images(str(pptx_file), output_dir)
        total_all_images += images
        total_all_slides += slides
    
    print("=" * 60)
    print(f"\nBatch Summary:")
    print(f"  Files processed: {len(pptx_files)}")
    print(f"  Total images extracted: {total_all_images}")
    print(f"  Total slides processed: {total_all_slides}")
    print()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("""\nImage Extraction Tool for Course Materials

Usage:
  Single file:  python extract_pptx_images.py <pptx_file> [output_dir]
  Batch mode:   python extract_pptx_images.py --batch [directory] [output_dir]

Examples:
  python extract_pptx_images.py "3. Class Title.pptx"
  python extract_pptx_images.py "3. Class Title.pptx" docs/classes/images
  python extract_pptx_images.py --batch . docs/classes/images
  python extract_pptx_images.py --batch /path/to/pptx/files

Output:
  Images are saved with naming pattern: class##_slide##_img##.ext
  Example: class03_slide05_img01.png
""")
        sys.exit(1)
    
    # Batch mode
    if sys.argv[1] == "--batch":
        directory = sys.argv[2] if len(sys.argv) > 2 else "."
        output_dir = sys.argv[3] if len(sys.argv) > 3 else "docs/classes/images"
        batch_extract(directory, output_dir)
    
    # Single file mode
    else:
        pptx_file = sys.argv[1]
        output_dir = sys.argv[2] if len(sys.argv) > 2 else "docs/classes/images"
        extract_images(pptx_file, output_dir)
